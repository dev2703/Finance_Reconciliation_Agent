from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from typing import Any

from packages.contracts import (
    BankTransaction,
    FinancialRecord,
    Invoice,
    LedgerEntry,
    Payment,
    Provenance,
    Settlement,
)

RECORD_TYPES = {
    "bank": BankTransaction,
    "bank_transaction": BankTransaction,
    "invoice": Invoice,
    "ledger": LedgerEntry,
    "ledger_entry": LedgerEntry,
    "payment": Payment,
    "settlement": Settlement,
}


class TabularParseError(ValueError):
    """Raised when a tabular row cannot be normalized."""


@dataclass
class RowParseError:
    row_number: int
    sheet: str | None
    field: str | None
    original_value: str | None
    message: str


@dataclass
class TabularParseResult:
    records: list[FinancialRecord] = field(default_factory=list)
    errors: list[RowParseError] = field(default_factory=list)


def normalize_header(value: str) -> str:
    return re.sub(r"[^a-z0-9]+", "_", value.strip().lower()).strip("_")


def parse_decimal(value: Any, *, field_name: str) -> Decimal:
    if value is None or str(value).strip() == "":
        raise TabularParseError(f"Missing required amount field: {field_name}")
    text = str(value).strip().replace(",", "").replace("$", "").replace("€", "").replace("£", "")
    negative = text.startswith("(") and text.endswith(")")
    if negative:
        text = text[1:-1]
    try:
        amount = Decimal(text)
    except InvalidOperation as exc:
        raise TabularParseError(f"Invalid decimal in {field_name}: {value!r}") from exc
    return -amount if negative else amount


def parse_date(value: Any, *, field_name: str) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value).strip() if value is not None else ""
    for format_string in ("%Y-%m-%d", "%m/%d/%Y", "%d/%m/%Y", "%Y/%m/%d"):
        try:
            return datetime.strptime(text, format_string).date()
        except ValueError:
            continue
    raise TabularParseError(f"Invalid date in {field_name}: {value!r}")


def _value(row: dict[str, Any], *names: str) -> Any:
    for name in names:
        if name in row and row[name] not in (None, ""):
            return row[name]
    return None


def normalize_row(
    row: dict[str, Any],
    *,
    record_type: str,
    source_name: str,
    row_number: int,
    sheet: str | None = None,
    column_mapping: dict[str, str] | None = None,
) -> FinancialRecord:
    normalized = {normalize_header(key): value for key, value in row.items() if key is not None}
    if column_mapping:
        mapped = {
            target: normalized.get(normalize_header(source), normalized.get(target))
            for target, source in column_mapping.items()
        }
        normalized = mapped | normalized

    model = RECORD_TYPES.get(record_type.lower().strip())
    if model is None:
        raise TabularParseError(f"Unsupported record type: {record_type}")

    amount = _value(normalized, "amount", "total", "value")
    if amount is None:
        debit = _value(normalized, "debit")
        credit = _value(normalized, "credit")
        if debit is not None or credit is not None:
            amount = parse_decimal(debit or credit, field_name="debit/credit")
            if credit is not None and debit is None:
                amount = -amount
    parsed: dict[str, Any] = {
        "external_id": _value(normalized, "external_id", "id", "transaction_id", "reference"),
        "amount": parse_decimal(amount, field_name="amount"),
        "currency": str(_value(normalized, "currency", "currency_code") or "USD").upper(),
        "record_date": parse_date(
            _value(normalized, "record_date", "transaction_date", "date", "posting_date"),
            field_name="date",
        ),
        "description": _value(normalized, "description", "memo", "narrative"),
        "provenance": Provenance(
            source_name=source_name,
            source_type="tabular",
            sheet=sheet,
            row_number=row_number,
            original_fields={
                str(key): "" if value is None else str(value) for key, value in row.items()
            },
            transformation_history=[
                "normalized column names",
                "parsed Decimal amount",
                "parsed date",
            ],
        ),
    }

    if model is Invoice:
        parsed.update(
            invoice_number=str(_value(normalized, "invoice_number", "invoice", "id") or ""),
            vendor_id=_value(normalized, "vendor_id", "vendor", "supplier_id"),
            due_date=(parse_date(_value(normalized, "due_date"), field_name="due_date")
                      if _value(normalized, "due_date") else None),
            tax_amount=(
                parse_decimal(_value(normalized, "tax_amount", "tax"), field_name="tax_amount")
                if _value(normalized, "tax_amount", "tax")
                else Decimal(0)
            ),
        )
    elif model is Payment:
        parsed["payment_reference"] = _value(normalized, "payment_reference", "payment_id")
    elif model is Settlement:
        parsed.update(
            settlement_reference=str(
                _value(normalized, "settlement_reference", "settlement_id", "id") or ""
            ),
            processor=_value(normalized, "processor"),
            fee_amount=(
                parse_decimal(_value(normalized, "fee_amount", "fee"), field_name="fee_amount")
                if _value(normalized, "fee_amount", "fee")
                else Decimal(0)
            ),
            settled_date=(parse_date(_value(normalized, "settled_date"), field_name="settled_date")
                          if _value(normalized, "settled_date") else None),
        )
    elif model is BankTransaction:
        parsed.update(
            account_id=str(_value(normalized, "account_id", "account") or "unknown"),
            transaction_type=str(_value(normalized, "transaction_type", "type") or "unknown"),
            bank_reference=_value(normalized, "bank_reference", "bank_id"),
            value_date=(parse_date(_value(normalized, "value_date"), field_name="value_date")
                        if _value(normalized, "value_date") else None),
        )
    elif model is LedgerEntry:
        parsed.update(
            journal_id=str(_value(normalized, "journal_id", "journal") or "unknown"),
            account_code=str(_value(normalized, "account_code", "account") or "unknown"),
            debit=(parse_decimal(_value(normalized, "debit"), field_name="debit")
                   if _value(normalized, "debit") else Decimal(0)),
            credit=(parse_decimal(_value(normalized, "credit"), field_name="credit")
                    if _value(normalized, "credit") else Decimal(0)),
            posting_date=(parse_date(_value(normalized, "posting_date"), field_name="posting_date")
                          if _value(normalized, "posting_date") else None),
        )
    return model.model_validate(parsed)


def parse_csv(
    content: bytes,
    *,
    record_type: str,
    source_name: str,
    column_mapping: dict[str, str] | None = None,
) -> list[FinancialRecord]:
    result = parse_csv_result(
        content,
        record_type=record_type,
        source_name=source_name,
        column_mapping=column_mapping,
    )
    if result.errors:
        raise TabularParseError(result.errors[0].message)
    return result.records


def parse_csv_result(
    content: bytes,
    *,
    record_type: str,
    source_name: str,
    column_mapping: dict[str, str] | None = None,
    max_rows: int | None = None,
) -> TabularParseResult:
    reader = csv.DictReader(io.StringIO(content.decode("utf-8-sig")))
    if not reader.fieldnames:
        raise TabularParseError("CSV has no header row")
    result = TabularParseResult()
    for index, row in enumerate(reader, start=2):
        if max_rows is not None and index - 1 > max_rows:
            result.errors.append(RowParseError(index, None, None, None, "CSV row limit exceeded"))
            break
        try:
            result.records.append(
                normalize_row(
                    row,
                    record_type=record_type,
                    source_name=source_name,
                    row_number=index,
                    column_mapping=column_mapping,
                )
            )
        except (TabularParseError, ValueError) as exc:
            result.errors.append(
                RowParseError(index, None, None, None, str(exc))
            )
    return result


def parse_xlsx(
    content: bytes,
    *,
    record_type: str,
    source_name: str,
    sheet_name: str | None = None,
    column_mapping: dict[str, str] | None = None,
) -> list[FinancialRecord]:
    result = parse_xlsx_result(
        content,
        record_type=record_type,
        source_name=source_name,
        sheet_name=sheet_name,
        column_mapping=column_mapping,
    )
    if result.errors:
        raise TabularParseError(result.errors[0].message)
    return result.records


def parse_xlsx_result(
    content: bytes,
    *,
    record_type: str,
    source_name: str,
    sheet_name: str | None = None,
    column_mapping: dict[str, str] | None = None,
    max_sheets: int | None = None,
    max_rows: int | None = None,
) -> TabularParseResult:
    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    if sheet_name and sheet_name not in workbook.sheetnames:
        raise TabularParseError(f"Unknown worksheet: {sheet_name}")
    sheets = [workbook[sheet_name]] if sheet_name else workbook.worksheets
    result = TabularParseResult()
    if max_sheets is not None and len(sheets) > max_sheets:
        result.errors.append(RowParseError(1, None, None, None, "XLSX sheet limit exceeded"))
        sheets = sheets[:max_sheets]
    try:
        for worksheet in sheets:
            rows = worksheet.iter_rows(values_only=True)
            headers = next(
                (list(row) for row in rows if any(value not in (None, "") for value in row)),
                None,
            )
            if not headers:
                continue
            header_names = [str(value or f"column_{index}") for index, value in enumerate(headers)]
            for row_number, values in enumerate(rows, start=2):
                if max_rows is not None and row_number - 1 > max_rows:
                    result.errors.append(
                        RowParseError(
                            row_number,
                            worksheet.title,
                            None,
                            None,
                            "XLSX row limit exceeded",
                        )
                    )
                    break
                if not any(value not in (None, "") for value in values):
                    continue
                row = dict(zip(header_names, values, strict=False))
                try:
                    result.records.append(
                        normalize_row(
                            row,
                            record_type=record_type,
                            source_name=source_name,
                            row_number=row_number,
                            sheet=worksheet.title,
                            column_mapping=column_mapping,
                        )
                    )
                except (TabularParseError, ValueError) as exc:
                    result.errors.append(
                        RowParseError(row_number, worksheet.title, None, None, str(exc))
                    )
    finally:
        workbook.close()
    return result


def preview_rows(content: bytes, *, file_type: str, limit: int = 25) -> list[Any]:
    if file_type == "csv":
        reader = csv.DictReader(io.StringIO(content.decode("utf-8-sig")))
        return [dict(row) for _, row in zip(range(limit), reader, strict=False)]
    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    try:
        worksheet = workbook.worksheets[0]
        rows = worksheet.iter_rows(values_only=True)
        return [list(row) for _, row in zip(range(limit), rows, strict=False)]
    finally:
        workbook.close()